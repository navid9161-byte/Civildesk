"""ورود دریافت و پرداخت‌های گذشته از Excel (داده‌ها ساختگی‌اند، با ساختار خروجی یک برنامه‌ی حسابداری شخصی)."""
import io

import openpyxl

from civildesk import db, importer


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ریز هزینه‌ها"
    ws.append(["تاریخ", "شرح هزینه", "مبلغ (تومان)", "دسته‌بندی", "خودرو", "توضیحات"])
    ws.append(["۱۴۰۵/۰۶/۰۱", "سرویس روغن", 5000000, "خودرو", "پراید", "فیلتر هوا"])
    ws.append(["۱۴۰۵/۰۶/۰۲", "نان", 250000, "خوراک", None, None])
    ws.append(["۱۴۰۵/۰۶/۰۲", "نان", 250000, "خوراک", None, None])  # دو خرید یکسان در یک روز: هر دو ثبت می‌شوند
    ws.append(["۱۴۰۵/۰۶/۰۳", "بنزین", 60000, "خودرو", "نامشخص", None])
    s2 = wb.create_sheet("خلاصه روزانه")
    s2.append(["تاریخ", "جمع هزینه (تومان)"])
    s2.append(["۱۴۰۵/۰۶/۰۱", 5000000])
    s3 = wb.create_sheet("خلاصه دسته‌بندی")
    s3.append(["دسته‌بندی", "جمع (تومان)"])
    s3.append(["خوراک", 500000])
    s4 = wb.create_sheet("تأمین مالی و بدهی")
    s4.append(["تاریخ", "شرح", "مبلغ (تومان)", "نوع", "توضیحات"])
    s4.append(["۱۴۰۵/۰۶/۰۵", "دریافت از خانواده", 3000000, "تأمین مالی دریافتی", None])
    s4.append(["تاریخ دقیق نامشخص؛ ماه قبل", "قرض قبلی", 1000000, "بدهی/قرض دریافتی", None])
    s4.append(["۱۴۰۵/۰۶/۰۶", "بازپرداخت بدهی", 500000, "بازپرداخت بدهی", None])
    s5 = wb.create_sheet("یادداشت و یادآوری")
    s5.append(["موضوع", "جزئیات"])
    s5.append(["یادآوری خودرو", "تعویض تسمه در پاییز"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_preview_commit_and_no_duplicates():
    data = _xlsx()
    r = importer.run(data, "هزینه‌ها.xlsx")  # پیش‌نمایش: چیزی ثبت نمی‌شود
    assert not r["committed"] and r["new"] == 7 and r["notes"] == 1
    types = {s["name"]: s["type"] for s in r["sheets"]}
    assert types["خلاصه روزانه"] == types["خلاصه دسته‌بندی"] == "skipped"  # جمع‌ها دو بار حساب نمی‌شوند
    assert r["expense"] == (5000000 + 250000 * 2 + 60000 + 500000) * 10  # تومان ← ریال
    assert r["income"] == 4000000 * 10
    with db.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0] == 0

    done = importer.run(data, "هزینه‌ها.xlsx", commit=True)
    assert done["new"] == 7
    with db.connect() as conn:
        rows = {r["description"]: dict(r) for r in conn.execute("SELECT * FROM transactions")}
        notes = conn.execute("SELECT title, content FROM notes").fetchall()
    oil = rows["سرویس روغن — خودرو: پراید — فیلتر هوا"]
    assert oil["amount"] == 50000000 and oil["category"] == "خودرو" and oil["tx_date"] == "1405/06/01"
    assert "بنزین" in rows and rows["بنزین"]["kind"] == "expense"  # «نامشخص» به شرح اضافه نمی‌شود
    loan = next(v for k, v in rows.items() if k.startswith("قرض قبلی"))
    assert loan["kind"] == "income" and loan["tx_date"] == "1405/05/31" and "تاریخ دقیق نامشخص" in loan["description"]
    assert rows["بازپرداخت بدهی"]["kind"] == "expense"
    assert [tuple(n) for n in notes] == [("یادآوری خودرو", "تعویض تسمه در پاییز")]

    again = importer.run(data, "هزینه‌ها.xlsx", commit=True)  # بارگذاری دوباره‌ی همان فایل
    assert again["new"] == 0 and again["duplicates"] == 7 and again["notes"] == 0


def test_import_csv_with_income_and_expense_columns():
    csv_text = "تاریخ,شرح,واریز (ریال),برداشت (ریال)\n1405/07/01,حقوق مهر,900000000,\n1405/07/02,اجاره,,300000000\n"
    r = importer.run(csv_text.encode("utf-8"), "bank.csv", commit=True)
    assert r["new"] == 2 and r["income"] == 900000000 and r["expense"] == 300000000


def test_finance_summary_by_month_and_category():
    from civildesk import services

    importer.run(_xlsx(), "هزینه‌ها.xlsx", commit=True)
    with db.connect() as conn:
        s = services.finance_summary(conn, "1405/06/01", "1405/06/30")
    assert s["count"] == 6 and s["first"] == "1405/06/01" and s["last"] == "1405/06/06"
    months = {(m["month"], m["kind"]): m["total"] for m in s["by_month"]}
    assert months[("1405/06", "income")] == 30000000
    cats = {(c["kind"], c["category"]): c["total"] for c in s["by_category"]}
    assert cats[("expense", "خوراک")] == 5000000
