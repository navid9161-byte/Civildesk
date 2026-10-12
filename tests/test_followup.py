"""پیگیری نامه‌ها، هشدارهای کنترل پروژه، پیش‌نویس‌های Word و یادگیری از اصلاح‌ها (داده‌ها ساختگی‌اند)."""
import io
import json

from civildesk import archive, db, documents, drafts, learning

L1 = """شماره: 1405/120/ص
تاریخ: 1405/05/10
جناب آقای مهندس احمدی
مدیرعامل محترم شرکت توسعه نمونه
موضوع: تحویل زمین بازه دوم
با سلام و احترام
احتراماً، با توجه به تاخیر در تحویل زمین بازه دوم و توقف عملیات خاکریزی، خواهشمند است دستور فرمایید زمین در اسرع وقت تحویل گردد.
"""
L2 = """شماره: 1405/233/م
تاریخ: 1405/05/20
جناب آقای مهندس رضایی
مدیرعامل محترم شرکت پیمانکاری نمونه
موضوع: پاسخ درباره تحویل زمین
با سلام و احترام
بازگشت به نامه شماره 1405/120/ص مورخ 1405/05/10، به استحضار می‌رساند زمین بازه دوم تا پایان ماه جاری تحویل خواهد شد.
"""
L3 = """شماره: 1405/150/ص
تاریخ: 1405/05/15
جناب آقای مهندس احمدی
مدیرعامل محترم شرکت توسعه نمونه
موضوع: عدم پرداخت صورت وضعیت
با سلام و احترام
پیرو نامه شماره 1405/99/ص، با توجه به عدم پرداخت صورت وضعیت شماره چهار، خواهشمند است دستور فرمایید نسبت به پرداخت آن اقدام گردد.
"""


def _project(**kw):
    with db.connect() as conn:
        return db.create(conn, "projects", {"name": "راه دسترسی نمونه", **kw})


def _add(text: str, name: str, pid: int | None) -> int:
    d = documents.add_document(io.BytesIO(text.encode()), name, project_id=pid)
    documents.worker.run_pending()
    return d["id"]


def test_letter_threads_and_unanswered():
    p = _project()
    a, b, c = (_add(t, f"{n}.txt", p["id"]) for t, n in ((L1, "l1"), (L2, "l2"), (L3, "l3")))
    reg = archive.letters_register(p["id"])
    rows = {x["id"]: x for x in reg["letters"]}
    assert [r["id"] for r in rows[b]["reply_to"]] == [a]  # «بازگشت به نامه …»
    assert [r["id"] for r in rows[a]["replies"]] == [b]
    assert rows[a]["waiting"] is None  # پاسخ گرفته
    assert rows[c]["needs_reply"] and rows[c]["waiting"] is not None and not rows[c]["replies"]
    assert rows[c]["missing"] == [{"no": "1405/99/ص", "date": None}]  # نامه‌ی ارجاعی در بایگانی نیست
    assert reg["unanswered"] == 1


def test_control_alerts_and_chart():
    p = _project(start_date="1404/05/01", end_date="1405/08/30", contract_amount=8_880_000_000_000)
    _add(L3, "l3.txt", p["id"])
    inv = _add("صورت وضعیت موقت شماره 5\nکارکرد", "وضعیت 5.txt", p["id"])
    info = {"invoice_no": {"value": "5"}, "work_total": {"value": 1_777_566_341_563, "page": 1},
            "claimed_period": {"value": 300_000_000_000}, "approved_period": {"value": 250_000_000_000}}
    with db.connect() as conn:
        conn.execute("UPDATE documents SET category='invoice', doc_date='1405/01/31', info=?, edited='[\"category\",\"doc_date\"]' "
                     "WHERE id=?", (json.dumps(info), inv))
        conn.execute("UPDATE documents SET doc_date='1405/03/01', edited='[\"doc_date\"]' WHERE id != ?", (inv,))
    s = archive.project_summary(p["id"])
    texts = " | ".join(a["text"] for a in s["alerts"])
    assert "عقب‌تر از برنامه" in texts  # زمان سپری‌شده خیلی بیشتر از کارکرد ۲۰٪
    assert "کمتر از درخواست پیمانکار" in texts  # تأیید ۲۵۰ از ۳۰۰
    assert "نامه‌ی درخواست‌دار" in texts and "روز است صورت‌وضعیت جدیدی" in texts
    assert s["alerts"][0]["level"] == "red"
    assert s["chart"]["points"][0]["value"] == 1_777_566_341_563 and s["chart"]["contract"] == 8_880_000_000_000


def test_word_drafts():
    import docx

    p = _project(start_date="1404/05/01", end_date="1405/08/30", employer="شرکت توسعه نمونه")
    _add(L1, "l1.txt", p["id"])
    _add(L3, "l3.txt", p["id"])
    name, data = drafts.build(p["id"], "extension")
    text = "\n".join(par.text for par in docx.Document(io.BytesIO(data)).paragraphs)
    assert name.startswith("درخواست تمدید") and "تحویل زمین" in text and "عدم پرداخت" in text
    assert "شماره‌ی ۱۴۰۵/۱۲۰/ص" in text  # مستند با شماره و تاریخ نامه
    name, data = drafts.build(p["id"], "monthly", "1405/05")
    d = docx.Document(io.BytesIO(data))
    assert "مرداد ۱۴۰۵" in d.paragraphs[0].text and len(d.tables[1].rows) == 3  # دو نامه‌ی مرداد
    try:
        drafts.build(p["id"], "invoice")
        raise AssertionError("بدون صورت‌وضعیت باید خطا بدهد")
    except db.ValidationError:
        pass


def test_learning_from_corrections():
    p = _project()
    t1 = "فرم درخواست مصالح کارگاه\nکد نامه: 7781-ع\nمورخه 1405/03/02\nشرح: سیمان و میلگرد\n"
    t2 = "فرم درخواست مصالح کارگاه\nکد نامه: 9902-ع\nمورخه 1405/04/11\nشرح: شن و ماسه\n"
    a = _add(t1, "a.txt", p["id"])
    b = _add(t2, "b.txt", p["id"])
    assert documents.get_document(b)["doc_no"] is None  # برچسب «کد نامه» را نمی‌شناسد
    assert documents.get_document(b)["category"] == "other"
    learned = archive.update_meta(a, {"doc_no": "7781-ع", "category": "order"})
    assert set(learned) == {"doc_no", "category"}
    rules = learning.load()
    assert rules["doc_no"][0]["label"] == "کد نامه"
    doc_b = documents.get_document(b)  # سند دیگر با الگوی یادگرفته دوباره خوانده شد
    assert doc_b["doc_no"] == "9902-ع" and doc_b["category"] == "order"
    assert archive.update_meta(a, {"subject": "فقط موضوع"}) == []  # موضوع چیزی یاد نمی‌دهد
