"""پیش‌نویس‌های آماده (Word) از روی اطلاعات بایگانی پروژه: نامه‌ی ارسال صورت‌وضعیت، درخواست تمدید مدت و گزارش ماهانه.

متن‌ها پیش‌نویس‌اند: جاهایی که اطلاعاتش در بایگانی نیست با «[…]» مشخص شده تا کاربر کامل کند.
"""
from __future__ import annotations

import io
import re
from typing import Any

from . import archive, db, jalali, textnorm

FONT = "B Nazanin"
KINDS = {"invoice": "نامه‌ی ارسال صورت‌وضعیت", "extension": "درخواست تمدید مدت", "monthly": "گزارش ماهانه"}
# نشانه‌های تأخیرِ ناشی از طرف مقابل (برای مستندات درخواست تمدید)
_DELAY = ("تاخیر", "تأخیر", "عدم تحویل", "عدم پرداخت", "عدم صدور", "عدم ابلاغ", "عدم تامین", "عدم تأمین", "کمبود",
          "توقف", "تعلیق", "دیرکرد", "مطالبات", "معارض", "عدم تکمیل", "عدم ارسال", "عدم تایید", "عدم تأیید", "بارندگی")


def fa(x: Any) -> str:
    return str(x).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")) if x is not None else ""


def money(n: int | None) -> str:
    return fa(f"{n:,}") + " ریال" if n else "[…] ریال"


# ───────────────────────── Word راست‌به‌چپ ─────────────────────────

# عناصری که در XML ورد باید بعد از bidi / szCs / bidiVisual بیایند (Word به ترتیب عناصر حساس است)
_PPR_AFTER_BIDI = ("w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                   "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                   "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_RPR_AFTER_SZCS = ("w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
                   "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
_TBLPR_AFTER_BIDI = ("w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
                     "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")


def _cs_size(rpr, half_points: int) -> None:
    """اندازه‌ی قلم متن فارسی (complex script) که Word جدا از اندازه‌ی لاتین نگه می‌دارد."""
    from docx.oxml.ns import qn

    old = rpr.find(qn("w:szCs"))
    if old is not None:
        rpr.remove(old)
    rpr.insert_element_before(rpr.makeelement(qn("w:szCs"), {qn("w:val"): str(half_points)}), *_RPR_AFTER_SZCS)


class Doc:
    def __init__(self) -> None:
        import docx
        from docx.shared import Cm, Pt

        self.d = docx.Document()
        sec = self.d.sections[0]
        sec.right_margin = sec.left_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2)
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(12)
        self._rtl_style(st)

    @staticmethod
    def _rtl_style(style) -> None:
        from docx.oxml.ns import qn

        rpr = style.element.get_or_add_rPr()
        fonts = rpr.get_or_add_rFonts()
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(k), FONT)
        _cs_size(rpr, 24)

    @staticmethod
    def _rtl(p, align: str = "right") -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        ppr = p._p.get_or_add_pPr()
        if ppr.find(qn("w:bidi")) is None:  # ترتیب عناصر طبق استاندارد Word
            ppr.insert_element_before(ppr.makeelement(qn("w:bidi"), {}), *_PPR_AFTER_BIDI)
        # پاراگراف راست‌به‌چپ خودش از راست شروع می‌شود؛ «right» صریح را Word در این حالت برعکس (چپ) نشان می‌دهد
        if align != "right":
            p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
        for r in p.runs:
            r.font.rtl = True
            if r.bold:
                r.font.cs_bold = True

    def para(self, text: str = "", bold: bool = False, size: int | None = None, align: str = "right",
             space: int = 4) -> None:
        from docx.shared import Pt

        p = self.d.add_paragraph()
        r = p.add_run(fa(text))
        r.bold = bold
        if size:
            r.font.size = Pt(size)
            _cs_size(r._r.get_or_add_rPr(), size * 2)
        p.paragraph_format.space_after = Pt(space)
        self._rtl(p, align)

    def heading(self, text: str) -> None:
        self.para(text, bold=True, size=14, space=6)

    def bullets(self, items: list[str], numbered: bool = True) -> None:
        for i, it in enumerate(items, 1):
            self.para(f"{i}- {it}" if numbered else f"• {it}", space=2)

    def table(self, header: list[str], rows: list[list[Any]]) -> None:
        from docx.oxml.ns import qn

        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        tblpr = t._tbl.tblPr
        tblpr.insert_element_before(tblpr.makeelement(qn("w:bidiVisual"), {}), *_TBLPR_AFTER_BIDI)
        for cell, h in zip(t.rows[0].cells, header):
            cell.text = ""
            r = cell.paragraphs[0].add_run(fa(h))
            r.bold = True
            self._rtl(cell.paragraphs[0], "center")
        for row in rows:
            cells = t.add_row().cells
            for cell, v in zip(cells, row):
                cell.text = ""
                cell.paragraphs[0].add_run(fa(v if v is not None else ""))
                self._rtl(cell.paragraphs[0])
        self.para("")

    def bytes(self) -> bytes:
        buf = io.BytesIO()
        self.d.save(buf)
        return buf.getvalue()


# ───────────────────────── داده‌ها ─────────────────────────


def _ctx(project_id: int) -> dict[str, Any]:
    s = archive.project_summary(project_id)
    if not s["project"]:
        raise db.NotFound("پروژه پیدا نشد")
    f = {x["key"]: x["value"] for x in s["facts"]}
    p = s["project"]
    for k in ("employer", "contractor", "consultant", "start_date", "end_date", "contract_amount"):
        f.setdefault(k, p.get(k))
    f.setdefault("contract_no", p.get("code"))
    s["f"] = f
    return s


def _consultant(name: str | None) -> str | None:
    if not name:
        return None
    return name if "مشاور" in name else f"مهندسین مشاور {name}"


def _header(doc: Doc, s: dict, to: str | None, subject: str) -> None:
    today = db.today_str()
    doc.para(f"تاریخ: {today}", align="right", space=0)
    doc.para("شماره: […]", space=0)
    doc.para("پیوست: […]", space=10)
    doc.para(f"{to or '[گیرنده]'}", bold=True, space=8)
    doc.para(f"موضوع: {subject}", bold=True, space=8)
    doc.para("با سلام و احترام", space=6)


def _sign(doc: Doc, s: dict) -> None:
    doc.para("")
    doc.para("با تشکر", align="center", space=0)
    doc.para(s["f"].get("contractor") or "[نام شرکت / امضا]", bold=True, align="center")


def _project_line(s: dict) -> str:
    f = s["f"]
    t = f"پروژه‌ی «{s['project']['name']}»"
    if f.get("contract_no"):
        t += f" (قرارداد شماره‌ی {f['contract_no']}" + (f" مورخ {f['contract_date']}" if f.get("contract_date") else "") + ")"
    return t


def invoice_letter(project_id: int, doc_id: int | None = None) -> tuple[str, bytes]:
    s = _ctx(project_id)
    invs = s["invoices"]
    inv = next((x for x in invs if x.get("doc_id") == doc_id), None) if doc_id else (invs[-1] if invs else None)
    if not inv:
        raise db.ValidationError("هنوز صورت‌وضعیتی در بایگانی این پروژه نیست")
    no = inv.get("no") or "[…]"
    doc = Doc()
    subj = f"ارسال صورت‌وضعیت {inv.get('kind_fa') or 'موقت'} شماره‌ی {no}" + (f" (دوره‌ی {inv['period']})" if inv.get("period") else "")
    _header(doc, s, _consultant(s["f"].get("consultant")), subj)
    claimed = inv.get("claimed_period") or inv.get("work_period")
    total = inv.get("claimed_total") or inv.get("work_total")
    body = (f"احتراماً، به پیوست صورت‌وضعیت {inv.get('kind_fa') or 'موقت'} شماره‌ی {no} {_project_line(s)}"
            + (f" مربوط به دوره‌ی کارکرد {inv['period']}" if inv.get("period") else "") + " جهت بررسی و تأیید ارسال می‌گردد.")
    doc.para(body, align="justify")
    rows = []
    if inv.get("work_prev"):
        rows.append(["کارکرد تجمعی تا دوره‌ی قبل", money(inv["work_prev"])])
    rows.append(["کارکرد این دوره (درخواستی)", money(claimed)])
    rows.append(["کارکرد تجمعی تا این دوره", money(total)])
    if s["f"].get("contract_amount") and total:
        rows.append(["نسبت به مبلغ قرارداد", f"{round(100 * total / s['f']['contract_amount'], 1)}٪"])
    doc.table(["شرح", "مبلغ"], rows)
    doc.para("خواهشمند است دستور فرمایید پس از بررسی، نسبت به تأیید و ارسال آن به کارفرما جهت پرداخت اقدام لازم معمول گردد.",
             align="justify")
    _sign(doc, s)
    doc.para(f"رونوشت: کارفرمای محترم{(' – ' + s['f']['employer']) if s['f'].get('employer') else ''} جهت استحضار", space=0)
    return f"ارسال صورت‌وضعیت {no} - {s['project']['name']}.docx", doc.bytes()


def delay_evidence(s: dict) -> list[str]:
    """مستندات تأخیر: نامه‌ها و گزارش‌هایی که از تأخیر، کمبود، عدم تحویل/پرداخت… گفته‌اند (قدیمی‌ترین اول)."""
    out = []
    for d in sorted(s["documents"], key=lambda d: d["date"]):
        if d["status"] != "ready":
            continue
        cat = d.get("category")
        if cat in ("letter", "order", "minutes"):
            summ = archive.letter_summary(d, 260) if cat != "minutes" else " ".join((d["info"].get("items") or {}).get("value") or [])[:260]
            norm = textnorm.normalize(summ)
            if summ and any(textnorm.normalize(k) in norm for k in _DELAY):
                kind = {"letter": "نامه", "order": "ابلاغیه", "minutes": "صورتجلسه"}[cat]
                ref = f"{kind}" + (f" شماره‌ی {d['doc_no']}" if d.get("doc_no") else "") + (f" مورخ {d['doc_date']}" if d.get("doc_date") else "")
                out.append(f"{ref}: {summ}")
        elif cat == "report":
            for p in (d["info"].get("problems") or {}).get("value") or []:
                out.append(f"گزارش ماهانه‌ی {(d['info'].get('period') or {}).get('value') or d['date']}: {p}")
    return out[:25]


def extension_request(project_id: int) -> tuple[str, bytes]:
    s = _ctx(project_id)
    f = s["f"]
    doc = Doc()
    _header(doc, s, f"کارفرمای محترم {f['employer']}" if f.get("employer") else None,
            f"درخواست تمدید مدت {_project_line(s)}")
    start, end, today = f.get("start_date"), f.get("end_date"), db.today_str()
    intro = f"احتراماً، به استحضار می‌رساند مدت {_project_line(s)}"
    if start and end:
        intro += f" از تاریخ {start} تا {end} می‌باشد"
        try:
            total, passed = jalali.days_between(start, end), jalali.days_between(start, today)
            intro += f" که تاکنون {passed:,} روز ({round(100 * passed / total)}٪) از آن سپری شده است"
        except ValueError:
            pass
    intro += "."
    doc.para(intro, align="justify")
    last = s["invoices"][-1] if s["invoices"] else None
    if last and last.get("work_total"):
        t = f"بر اساس آخرین صورت‌وضعیت (شماره‌ی {last.get('no') or '…'})، کارکرد تجمعی {money(last['work_total'])}"
        if f.get("contract_amount"):
            t += f" معادل {round(100 * last['work_total'] / f['contract_amount'], 1)}٪ مبلغ قرارداد"
        doc.para(t + " است.", align="justify")
    ev = delay_evidence(s)
    doc.para("در طول اجرای کار، عوامل زیر که خارج از اختیار این پیمانکار بوده موجب تأخیر در پیشرفت عملیات شده و مراتب "
             "طی مکاتبات زیر به اطلاع رسیده است:", align="justify")
    doc.bullets(ev or ["[عامل تأخیر و شماره و تاریخ نامه‌ی مربوط]"])
    doc.para("با توجه به موارد فوق و مستند به شرایط عمومی پیمان، خواهشمند است دستور فرمایید مدت قرارداد به میزان "
             "[…] ماه تمدید و نتیجه به این شرکت ابلاغ گردد. محاسبه‌ی تفصیلی تأخیرات مجاز به پیوست تقدیم می‌گردد.",
             align="justify")
    _sign(doc, s)
    if f.get("consultant"):
        doc.para(f"رونوشت: {_consultant(f['consultant'])} جهت استحضار و اعلام نظر", space=0)
    return f"درخواست تمدید - {s['project']['name']}.docx", doc.bytes()


def _month_of(d: str | None) -> str:
    return (d or "")[:7]


def monthly_report(project_id: int, month: str | None = None) -> tuple[str, bytes]:
    s = _ctx(project_id)
    f = s["f"]
    month = month if month and re.fullmatch(r"1[34]\d\d/\d\d", month) else _month_of(jalali.add_days(db.today_str(), -15))
    y, m = (int(x) for x in month.split("/"))
    month_fa = f"{jalali.MONTHS_FA[m - 1]} {y}"
    end = archive.month_end(y, m)
    docs = [d for d in s["documents"] if d["status"] == "ready"]
    in_month = [d for d in docs if _month_of(d.get("doc_date")) == month]
    report = next((d for d in sorted(docs, key=lambda d: d["date"], reverse=True)
                   if d.get("category") == "report" and _month_of(d.get("doc_date")) == month), None)
    doc = Doc()
    doc.para(f"گزارش ماهانه‌ی پیشرفت کار – {month_fa}", bold=True, size=16, align="center", space=2)
    doc.para(s["project"]["name"], bold=True, size=13, align="center", space=12)

    doc.heading("۱- مشخصات پروژه")
    rows = [[label, f.get(key)] for key, label in (
        ("employer", "کارفرما"), ("consultant", "مشاور"), ("contractor", "پیمانکار"), ("contract_no", "شماره‌ی قرارداد"),
        ("start_date", "تاریخ شروع"), ("end_date", "تاریخ پایان")) if f.get(key)]
    if f.get("contract_amount"):
        rows.append(["مبلغ قرارداد", money(f["contract_amount"])])
    doc.table(["عنوان", "شرح"], rows or [["[…]", "[…]"]])

    doc.heading("۲- وضعیت زمانی و مالی")
    lines = []
    if f.get("start_date") and f.get("end_date"):
        try:
            total, passed = jalali.days_between(f["start_date"], f["end_date"]), jalali.days_between(f["start_date"], end)
            lines.append(f"تا پایان {month_fa}، {passed:,} روز از {total:,} روز مدت قرارداد ({round(100 * passed / total)}٪) سپری شده است.")
        except ValueError:
            pass
    invs = [x for x in s["invoices"] if (x.get("date") or "") <= end]
    if invs and invs[-1].get("work_total"):
        last = invs[-1]
        t = f"آخرین صورت‌وضعیت (شماره‌ی {last.get('no') or '…'}): کارکرد تجمعی {money(last['work_total'])}"
        if f.get("contract_amount"):
            t += f" معادل {round(100 * last['work_total'] / f['contract_amount'], 1)}٪ مبلغ قرارداد"
        lines.append(t + ".")
    doc.bullets(lines or ["[وضعیت زمانی و مالی]"], numbered=False)

    def section(title: str, key: str, fallback: list[str]) -> None:
        doc.heading(title)
        items = ((report or {}).get("info", {}).get(key) or {}).get("value") or fallback
        doc.bullets(items or ["[…]"])

    section("۳- اقدامات انجام‌شده در این ماه", "done", [])
    letters = [d for d in in_month if d.get("category") in ("letter", "order")]
    doc.heading(f"۴- مکاتبات ماه ({fa(len(letters))} مورد)")
    if letters:
        doc.table(["تاریخ", "شماره", "موضوع", "خلاصه"],
                  [[d.get("doc_date"), d.get("doc_no") or "", (d.get("subject") or d["title"])[:80],
                    archive.letter_summary(d, 160)] for d in sorted(letters, key=lambda d: d["date"])])
    else:
        doc.para("مکاتبه‌ای از این ماه در بایگانی ثبت نشده است.")
    minutes = [d for d in in_month if d.get("category") == "minutes"]
    if minutes:
        doc.heading("۵- صورتجلسات")
        for d in minutes:
            doc.para(f"صورتجلسه‌ی {d.get('doc_date')}: {d.get('subject') or d['title']}", bold=True, space=2)
            doc.bullets((d["info"].get("items") or {}).get("value") or [])
    # هشدارهای امروز فقط برای گزارش ماه جاری/گذشته معنا دارند، نه برای ماه‌های قدیمی‌تر
    recent = month >= _month_of(jalali.add_days(db.today_str(), -45))
    probs = [a["text"] for a in s.get("alerts", []) if a["level"] in ("red", "amber")] if recent else []
    section("۶- مشکلات و موانع", "problems", probs)
    nxt_month = _month_of(jalali.add_days(end, 2))
    upcoming = [f"{d['due']}: {d['text'][:150]}" for d in s["deadlines"] if _month_of(d["due"]) == nxt_month]
    section("۷- برنامه‌ی ماه آینده", "next", upcoming)
    doc.para("")
    doc.para("تهیه‌کننده: […]          تأیید: […]", align="center")
    return f"گزارش ماهانه {month_fa} - {s['project']['name']}.docx", doc.bytes()


def build(project_id: int, kind: str, month: str | None = None, doc_id: int | None = None) -> tuple[str, bytes]:
    if kind == "invoice":
        return invoice_letter(project_id, doc_id)
    if kind == "extension":
        return extension_request(project_id)
    if kind == "monthly":
        return monthly_report(project_id, month)
    raise db.NotFound("نوع پیش‌نویس نامعتبر است")
